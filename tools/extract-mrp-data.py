#!/usr/bin/env python3
"""Extract sheet DATA columns A-U from Material_M&P MRP 4-1_Ver_(Toy).xlsm.

Writes a UTF-8 CSV whose first line is the Excel header row (row 2) verbatim and
whose remaining lines are the data rows from row 3 down, in sheet order. Cell
values are the stored values, not the formatted display text, so precision is
kept exactly as Excel holds it.

    python3 tools/extract-mrp-data.py "Material_M&P MRP 4-1_Ver_(Toy).xlsm" mrp-data.csv

Feed the CSV to mrp-import.html. The workbook and the CSV both hold real company
data and must stay out of this repository.
"""
import csv, sys, warnings

try:
    import openpyxl
except ImportError:
    sys.exit("ต้องติดตั้ง openpyxl ก่อน: pip install openpyxl")

warnings.filterwarnings("ignore")
COLUMNS = 21  # A..U


def main(source, target):
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if "DATA" not in book.sheetnames:
        sys.exit(f"ไม่พบชีต DATA ใน {source}")
    sheet = book["DATA"]
    header, rows = None, []
    for index, row in enumerate(sheet.iter_rows(min_col=1, max_col=COLUMNS, values_only=True), 1):
        if index == 1:
            continue                      # row 1 holds the lookup helpers, not headers
        if index == 2:
            header = ["" if v is None else str(v) for v in row]
            continue
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            continue                      # trailing blank rows
        rows.append(["" if v is None else v for v in row])
    book.close()

    if not header or len(header) != COLUMNS:
        sys.exit("แถวหัวข้อ (แถวที่ 2) ไม่ครบ 21 คอลัมน์ A-U")

    with open(target, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)
    print(f"เขียน {len(rows):,} แถว x {COLUMNS} คอลัมน์ ลง {target}")
    print("หัวข้อ:", " | ".join(header))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
